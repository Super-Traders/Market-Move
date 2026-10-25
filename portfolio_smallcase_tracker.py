#!/usr/bin/env python3
"""Maintain a monthly Smallcase-style, equal-weight 20-stock portfolio in Excel."""

from __future__ import annotations

import argparse
import calendar
import math
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from config import DB_PATH
from portfolio_strategy_backtest_rank import (
    compute_ranked_stock_details,
    load_cache_symbols,
)

PORTFOLIO_COLUMNS = [
    "Symbol",
    "Quantity",
    "Entry_Date",
    "Entry_Price",
    "Cost_Basis",
    "Last_Price",
    "Last_Price_Date",
    "Rank",
    "Market_Value",
    "Unrealized_PnL",
]
SUMMARY_COLUMNS = [
    "Rebalance_Date",
    "Execution_Date",
    "Ranked_Stocks",
    "Holdings_Count",
    "Buys",
    "Sells",
    "Realized_PnL",
    "Opening_Portfolio_Value",
    "Closing_Portfolio_Value",
    "Monthly_Profit_Loss",
    "Monthly_Gain_%",
    "Cumulative_Profit_Loss",
    "Cumulative_Return_%",
    "CAGR_%",
    "Cash_Balance",
]
RANK_COLUMNS = [
    "Rebalance_Date",
    "Execution_Date",
    "Rank",
    "Symbol",
    "Price",
    "Price_Date",
    "Composite Score",
    "RRG Quadrant",
    "RRG Direction",
    "RS Percentile",
    "Above 200-DMA",
    "% vs 200-DMA",
    "Previous_Rank",
    "Status",
]
TRADE_COLUMNS = [
    "Rebalance_Date",
    "Execution_Date",
    "Symbol",
    "Action",
    "Rank",
    "Quantity",
    "Price",
    "Gross_Amount",
    "Cost_Basis",
    "Realized_PnL",
]
BUY_SELL_COLUMNS = [
    "Symbol",
    "Buy_Date",
    "Buy_Quantity",
    "Buy_Price",
    "Buy_Total",
    "Sell_Date",
    "Sell_Quantity",
    "Sell_Price",
    "Sell_Total",
    "Realized_PnL",
    "Open_Quantity",
    "Current_Value",
    "Unrealized_PnL",
    "Status",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Track a monthly RRG-ranked portfolio and append details to an Excel workbook.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("--db-path", default=str(DB_PATH), help="SQLite market-data cache.")
    parser.add_argument(
        "--start-date",
        type=date.fromisoformat,
        default=date(2026, 9, 25),
        metavar="YYYY-MM-DD",
        help="First scheduled investment date.",
    )
    parser.add_argument(
        "--as-of-date",
        type=date.fromisoformat,
        default=date.today(),
        metavar="YYYY-MM-DD",
        help="Process scheduled 25th-of-month rebalance dates through this date.",
    )
    parser.add_argument("--top-n", type=int, default=20, help="Number of target stocks.")
    parser.add_argument(
        "--exit-rank",
        type=int,
        default=60,
        help="Sell an owned stock when its rank is worse than this rank.",
    )
    parser.add_argument(
        "--initial-capital",
        type=float,
        default=1_000_000.0,
        help="Initial capital in rupees.",
    )
    parser.add_argument(
        "--excel-output",
        type=Path,
        default=Path("./smallcase_portfolio_tracker.xlsx"),
        help="Excel workbook used both for the report and persistent portfolio state.",
    )
    parser.add_argument(
        "--reset-workbook",
        action="store_true",
        help="Ignore existing workbook state and rebuild from --start-date.",
    )
    return parser.parse_args()


def monthly_rebalance_dates(
    start_date: date, through_date: date, after_date: date | None = None
) -> list[date]:
    """Generate monthly dates on the 25th, preserving the start date's day."""
    dates = []
    year, month = start_date.year, start_date.month
    while (year, month) <= (through_date.year, through_date.month):
        day = min(start_date.day, calendar.monthrange(year, month)[1])
        scheduled_date = date(year, month, day)
        if (
            scheduled_date >= start_date
            and scheduled_date <= through_date
            and (after_date is None or scheduled_date > after_date)
        ):
            dates.append(scheduled_date)
        month += 1
        if month > 12:
            month = 1
            year += 1
    return dates


def _load_workbook_sheets(path: Path) -> dict[str, pd.DataFrame]:
    if not path.exists():
        return {}
    with pd.ExcelFile(path, engine="openpyxl") as workbook:
        return {
            name: pd.read_excel(workbook, sheet_name=name)
            for name in workbook.sheet_names
        }


def _read_metadata(sheets: dict[str, pd.DataFrame]) -> dict[str, str]:
    frame = sheets.get("Metadata")
    if frame is None or frame.empty or not {"Key", "Value"}.issubset(frame.columns):
        if sheets:
            raise ValueError(
                "The Excel file exists but has no valid tracker Metadata sheet. "
                "Choose a new --excel-output path or restore the tracker workbook."
            )
        return {}
    return {
        str(row["Key"]): str(row["Value"])
        for _, row in frame.iterrows()
        if pd.notna(row["Key"]) and pd.notna(row["Value"])
    }


def _check_settings(
    metadata: dict[str, str],
    start_date: date,
    top_n: int,
    exit_rank: int,
    initial_capital: float,
) -> None:
    expected = {
        "Start_Date": start_date.isoformat(),
        "Top_N": str(top_n),
        "Exit_Rank": str(exit_rank),
    }
    for key, value in expected.items():
        if key in metadata and metadata[key] != value:
            raise ValueError(
                f"Tracker setting {key} is {metadata[key]}, not {value}. "
                "Use the original settings or a different workbook."
            )
    if "Initial_Capital" in metadata and not math.isclose(
        float(metadata["Initial_Capital"]), initial_capital, rel_tol=0, abs_tol=0.005
    ):
        raise ValueError(
            "Initial capital does not match the existing workbook. "
            "Use the original amount or a different workbook."
        )


def _quote(
    price_lookup: dict[str, pd.Series], symbol: str, as_of_date: date
) -> tuple[float | None, str | None]:
    prices = price_lookup.get(symbol)
    if prices is None or prices.empty:
        return None, None
    available = prices.loc[:pd.Timestamp(as_of_date)].dropna()
    if available.empty:
        return None, None
    return float(available.iloc[-1]), available.index[-1].date().isoformat()


def _position_map(frame: pd.DataFrame | None) -> dict[str, dict]:
    if frame is None or frame.empty:
        return {}
    positions = {}
    for _, row in frame.iterrows():
        symbol = str(row["Symbol"])
        positions[symbol] = {
            "Quantity": int(row["Quantity"]),
            "Entry_Date": str(row["Entry_Date"]),
            "Entry_Price": float(row["Entry_Price"]),
            "Cost_Basis": float(row["Cost_Basis"]),
            "Last_Price": float(row["Last_Price"]),
            "Last_Price_Date": str(row["Last_Price_Date"]),
            "Rank": int(row["Rank"]) if pd.notna(row["Rank"]) else None,
        }
    return positions


def _metadata_frame(metadata: dict[str, str]) -> pd.DataFrame:
    return pd.DataFrame(
        [{"Key": key, "Value": value} for key, value in metadata.items()],
        columns=["Key", "Value"],
    )


def _calculate_cagr(
    initial_capital: float, portfolio_value: float, start_date: str, end_date: str
) -> float:
    """Annualize portfolio growth over actual elapsed calendar days."""
    elapsed_days = (
        date.fromisoformat(end_date) - date.fromisoformat(start_date)
    ).days
    if elapsed_days <= 0:
        return 0.0
    if initial_capital <= 0 or portfolio_value <= 0:
        return float("nan")
    return ((portfolio_value / initial_capital) ** (365.25 / elapsed_days) - 1) * 100


def _enrich_monthly_summary(
    summary: pd.DataFrame, initial_capital: float
) -> pd.DataFrame:
    """Add monthly percentage gain and annualized CAGR, including to older workbooks."""
    summary = summary.copy()
    if summary.empty:
        return summary.reindex(columns=SUMMARY_COLUMNS)

    opening_values = pd.to_numeric(
        summary.get("Opening_Portfolio_Value", pd.Series(index=summary.index, dtype=float)),
        errors="coerce",
    )
    monthly_profit = pd.to_numeric(
        summary.get("Monthly_Profit_Loss", pd.Series(index=summary.index, dtype=float)),
        errors="coerce",
    )
    summary["Monthly_Gain_%"] = (
        monthly_profit.div(opening_values.where(opening_values != 0)) * 100
    )
    first_execution_date = str(summary.iloc[0]["Execution_Date"])
    summary["CAGR_%"] = [
        _calculate_cagr(
            initial_capital,
            float(value),
            first_execution_date,
            str(execution_date),
        )
        for value, execution_date in zip(
            pd.to_numeric(summary["Closing_Portfolio_Value"], errors="coerce"),
            summary["Execution_Date"],
        )
    ]
    return summary.reindex(columns=SUMMARY_COLUMNS)


def _previous_ranks(
    rank_history: pd.DataFrame | None, last_rebalance: str | None
) -> dict[str, int]:
    if rank_history is None or rank_history.empty or not last_rebalance:
        return {}
    last = rank_history.loc[rank_history["Rebalance_Date"].astype(str) == last_rebalance]
    return {
        str(row["Symbol"]): int(row["Rank"])
        for _, row in last.iterrows()
        if pd.notna(row.get("Rank"))
    }


def _run_rebalance(
    ranked_rows: list[dict],
    positions: dict[str, dict],
    cash_balance: float,
    initial_capital: float,
    top_n: int,
    exit_rank: int,
    rebalance_date: date,
    execution_date: date,
    price_lookup: dict[str, pd.Series],
    initial: bool,
) -> tuple[dict[str, dict], float, list[dict], dict[str, str]]:
    rank_by_symbol = {row["Symbol"]: int(row["Rank"]) for row in ranked_rows}
    price_by_symbol = {}
    price_date_by_symbol = {}
    for row in ranked_rows:
        price, price_date = _quote(price_lookup, row["Symbol"], execution_date)
        if price is not None and price > 0:
            price_by_symbol[row["Symbol"]] = price
            price_date_by_symbol[row["Symbol"]] = price_date
    trades = []
    statuses: dict[str, str] = {}
    starting_positions = set(positions)
    monthly_sale_proceeds = 0.0

    if initial:
        cash_balance = initial_capital
        per_stock_budget = initial_capital / top_n
        for row in ranked_rows:
            if len(positions) >= top_n:
                break
            symbol = row["Symbol"]
            price = price_by_symbol.get(symbol)
            if price is None or price > per_stock_budget:
                statuses[symbol] = "SKIPPED_UNAFFORDABLE"
                continue
            quantity = int(per_stock_budget // price)
            if quantity < 1:
                statuses[symbol] = "SKIPPED_UNAFFORDABLE"
                continue
            cost = quantity * price
            positions[symbol] = {
                "Quantity": quantity,
                "Entry_Date": execution_date.isoformat(),
                "Entry_Price": price,
                "Cost_Basis": cost,
                "Last_Price": price,
                "Last_Price_Date": price_date_by_symbol[symbol],
                "Rank": int(row["Rank"]),
            }
            cash_balance -= cost
            statuses[symbol] = "INITIAL_BUY"
            trades.append({
                "Rebalance_Date": rebalance_date.isoformat(),
                "Execution_Date": execution_date.isoformat(),
                "Symbol": symbol,
                "Action": "BUY",
                "Rank": int(row["Rank"]),
                "Quantity": quantity,
                "Price": price,
                "Gross_Amount": cost,
                "Cost_Basis": cost,
                "Realized_PnL": 0.0,
            })
        if len(positions) != top_n:
            raise ValueError(
                f"Only {len(positions)} affordable, rankable stocks were found; "
                f"{top_n} are required. Check cached prices or increase initial capital."
            )
    else:
        for symbol in list(positions):
            rank = rank_by_symbol.get(symbol, exit_rank + 1)
            if rank <= exit_rank:
                statuses[symbol] = "RETAINED"
                continue
            price, price_date = _quote(price_lookup, symbol, execution_date)
            if price is None or price <= 0:
                statuses[symbol] = "EXIT_PENDING_UNPRICED"
                continue
            position = positions.pop(symbol)
            quantity = int(position["Quantity"])
            proceeds = quantity * price
            basis = float(position["Cost_Basis"])
            pnl = proceeds - basis
            cash_balance += proceeds
            monthly_sale_proceeds += proceeds
            status = "SOLD_RANK_BELOW_60" if symbol in rank_by_symbol else "SOLD_UNRANKED"
            statuses[symbol] = status
            trades.append({
                "Rebalance_Date": rebalance_date.isoformat(),
                "Execution_Date": price_date or execution_date.isoformat(),
                "Symbol": symbol,
                "Action": "SELL",
                "Rank": rank_by_symbol.get(symbol, np.nan),
                "Quantity": quantity,
                "Price": price,
                "Gross_Amount": proceeds,
                "Cost_Basis": basis,
                "Realized_PnL": pnl,
            })

        slots = max(0, top_n - len(positions))
        if slots and monthly_sale_proceeds > 0:
            per_stock_budget = monthly_sale_proceeds / slots
            for row in ranked_rows:
                if slots <= 0:
                    break
                symbol = row["Symbol"]
                if symbol in positions:
                    continue
                if int(row["Rank"]) > exit_rank:
                    continue
                price = price_by_symbol.get(symbol)
                if price is None or price > per_stock_budget:
                    statuses[symbol] = "SKIPPED_UNAFFORDABLE"
                    continue
                quantity = int(min(per_stock_budget, monthly_sale_proceeds) // price)
                if quantity < 1:
                    statuses[symbol] = "SKIPPED_UNAFFORDABLE"
                    continue
                cost = quantity * price
                if cost > monthly_sale_proceeds + 1e-8 or cost > cash_balance + 1e-8:
                    continue
                positions[symbol] = {
                    "Quantity": quantity,
                    "Entry_Date": execution_date.isoformat(),
                    "Entry_Price": price,
                    "Cost_Basis": cost,
                    "Last_Price": price,
                    "Last_Price_Date": price_date_by_symbol[symbol],
                    "Rank": int(row["Rank"]),
                }
                cash_balance -= cost
                monthly_sale_proceeds -= cost
                statuses[symbol] = "NEW_BUY"
                trades.append({
                    "Rebalance_Date": rebalance_date.isoformat(),
                    "Execution_Date": execution_date.isoformat(),
                    "Symbol": symbol,
                    "Action": "BUY",
                    "Rank": int(row["Rank"]),
                    "Quantity": quantity,
                    "Price": price,
                    "Gross_Amount": cost,
                    "Cost_Basis": cost,
                    "Realized_PnL": 0.0,
                })
                slots -= 1

    for symbol, position in positions.items():
        price, price_date = _quote(price_lookup, symbol, execution_date)
        if price is not None and price > 0:
            position["Last_Price"] = price
            position["Last_Price_Date"] = price_date
        position["Rank"] = rank_by_symbol.get(symbol)
        if symbol not in statuses:
            statuses[symbol] = "RETAINED"

    for row in ranked_rows:
        statuses.setdefault(row["Symbol"], "UNOWNED")
    for symbol in starting_positions - set(rank_by_symbol):
        statuses.setdefault(symbol, "SOLD_UNRANKED")

    return positions, cash_balance, trades, statuses


def _build_rank_rows(
    ranked_rows: list[dict],
    statuses: dict[str, str],
    previous_ranks: dict[str, int],
    rebalance_date: date,
    execution_date: date,
) -> list[dict]:
    result = []
    ranked_symbols = set()
    for row in ranked_rows:
        symbol = row["Symbol"]
        ranked_symbols.add(symbol)
        result.append({
            "Rebalance_Date": rebalance_date.isoformat(),
            "Execution_Date": execution_date.isoformat(),
            **{key: row.get(key) for key in RANK_COLUMNS[2:12]},
            "Previous_Rank": previous_ranks.get(symbol),
            "Status": statuses.get(symbol, "UNOWNED"),
        })
    for symbol, status in statuses.items():
        if symbol in ranked_symbols:
            continue
        result.append({
            "Rebalance_Date": rebalance_date.isoformat(),
            "Execution_Date": execution_date.isoformat(),
            "Rank": None,
            "Symbol": symbol,
            "Price": None,
            "Price_Date": None,
            "Composite Score": None,
            "RRG Quadrant": None,
            "RRG Direction": None,
            "RS Percentile": None,
            "Above 200-DMA": None,
            "% vs 200-DMA": None,
            "Previous_Rank": previous_ranks.get(symbol),
            "Status": status,
        })
    return result


def _portfolio_frame(
    positions: dict[str, dict], price_lookup: dict[str, pd.Series], as_of_date: date
) -> pd.DataFrame:
    rows = []
    for symbol, position in sorted(positions.items()):
        price, price_date = _quote(price_lookup, symbol, as_of_date)
        if price is None or price <= 0:
            price = float(position["Last_Price"])
            price_date = position["Last_Price_Date"]
        market_value = int(position["Quantity"]) * price
        rows.append({
            "Symbol": symbol,
            "Quantity": int(position["Quantity"]),
            "Entry_Date": position["Entry_Date"],
            "Entry_Price": float(position["Entry_Price"]),
            "Cost_Basis": float(position["Cost_Basis"]),
            "Last_Price": price,
            "Last_Price_Date": price_date,
            "Rank": position.get("Rank"),
            "Market_Value": market_value,
            "Unrealized_PnL": market_value - float(position["Cost_Basis"]),
        })
    return pd.DataFrame(rows, columns=PORTFOLIO_COLUMNS)


def _build_buy_sell_ledger(
    trade_rows: list[dict], current_portfolio: pd.DataFrame
) -> pd.DataFrame:
    """Pair sell transactions with prior buys by symbol using FIFO lot matching."""
    rows = []
    buys_by_symbol: dict[str, list[dict]] = {}
    transactions = sorted(
        trade_rows,
        key=lambda trade: (
            str(trade.get("Execution_Date", trade.get("Rebalance_Date", ""))),
            0 if trade.get("Action") == "BUY" else 1,
        ),
    )
    for trade in transactions:
        symbol = str(trade["Symbol"])
        quantity = int(trade["Quantity"])
        if quantity <= 0:
            raise ValueError(f"Invalid non-positive trade quantity for {symbol}.")
        if trade["Action"] == "BUY":
            buy_total = float(trade.get("Gross_Amount", 0.0))
            buys_by_symbol.setdefault(symbol, []).append({
                "date": str(trade.get("Execution_Date", trade.get("Rebalance_Date", ""))),
                "price": float(trade["Price"]),
                "total": buy_total,
                "quantity": quantity,
                "remaining": quantity,
            })
        elif trade["Action"] == "SELL":
            remaining_to_match = quantity
            sell_total = float(trade.get("Gross_Amount", 0.0))
            sell_unit_price = float(trade["Price"])
            lots = buys_by_symbol.get(symbol, [])
            for lot in lots:
                matched_quantity = min(remaining_to_match, lot["remaining"])
                if matched_quantity <= 0:
                    continue
                buy_total = lot["total"] * matched_quantity / lot["quantity"]
                proceeds = sell_total * matched_quantity / quantity
                rows.append({
                    "Symbol": symbol,
                    "Buy_Date": lot["date"],
                    "Buy_Quantity": matched_quantity,
                    "Buy_Price": lot["price"],
                    "Buy_Total": buy_total,
                    "Sell_Date": str(trade.get("Execution_Date", trade.get("Rebalance_Date", ""))),
                    "Sell_Quantity": matched_quantity,
                    "Sell_Price": sell_unit_price,
                    "Sell_Total": proceeds,
                    "Realized_PnL": proceeds - buy_total,
                    "Open_Quantity": 0,
                    "Current_Value": None,
                    "Unrealized_PnL": None,
                    "Status": "CLOSED",
                })
                lot["remaining"] -= matched_quantity
                remaining_to_match -= matched_quantity
                if remaining_to_match == 0:
                    break
            if remaining_to_match:
                raise ValueError(
                    f"Sell ledger has {remaining_to_match} unmatched shares for {symbol}; "
                    "the trade history is incomplete or inconsistent."
                )
        else:
            raise ValueError(f"Unsupported trade action: {trade['Action']}")

    portfolio_by_symbol = {}
    if not current_portfolio.empty:
        portfolio_by_symbol = current_portfolio.set_index("Symbol").to_dict("index")
    open_quantities: dict[str, int] = {}
    for symbol, lots in buys_by_symbol.items():
        for lot in lots:
            remaining = int(lot["remaining"])
            if remaining <= 0:
                continue
            open_quantities[symbol] = open_quantities.get(symbol, 0) + remaining
            current = portfolio_by_symbol.get(symbol)
            current_value = None
            unrealized_pnl = None
            if current is not None:
                current_value = float(current["Last_Price"]) * remaining
                remaining_basis = float(lot["total"]) * remaining / lot["quantity"]
                unrealized_pnl = current_value - remaining_basis
            remaining_basis = float(lot["total"]) * remaining / lot["quantity"]
            rows.append({
                "Symbol": symbol,
                "Buy_Date": lot["date"],
                "Buy_Quantity": remaining,
                "Buy_Price": lot["price"],
                "Buy_Total": remaining_basis,
                "Sell_Date": None,
                "Sell_Quantity": None,
                "Sell_Price": None,
                "Sell_Total": None,
                "Realized_PnL": None,
                "Open_Quantity": remaining,
                "Current_Value": current_value,
                "Unrealized_PnL": unrealized_pnl,
                "Status": "OPEN",
            })

    for symbol, current in portfolio_by_symbol.items():
        if open_quantities.get(symbol, 0) != int(current["Quantity"]):
            raise ValueError(
                f"Trade history has {open_quantities.get(symbol, 0)} open shares of {symbol}, "
                f"but Current_Portfolio has {int(current['Quantity'])}."
            )

    return pd.DataFrame(rows, columns=BUY_SELL_COLUMNS)


def _write_workbook(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, index=False, sheet_name=name)

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="17365D")
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        if worksheet.max_row >= 1 and worksheet.max_column >= 1:
            worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
        for column_cells in worksheet.columns:
            letter = column_cells[0].column_letter
            width = max(
                (len(str(cell.value)) for cell in column_cells if cell.value is not None),
                default=10,
            )
            worksheet.column_dimensions[letter].width = min(max(width + 2, 12), 34)
    workbook.save(path)


def run_tracker(
    db_path: str | Path = DB_PATH,
    start_date: date = date(2026, 9, 25),
    as_of_date: date | None = None,
    top_n: int = 20,
    exit_rank: int = 60,
    initial_capital: float = 1_000_000.0,
    excel_output: str | Path = "./smallcase_portfolio_tracker.xlsx",
    reset_workbook: bool = False,
) -> dict:
    as_of_date = as_of_date or date.today()
    if top_n <= 0 or exit_rank < top_n or initial_capital <= 0:
        raise ValueError("top_n and initial_capital must be positive; exit_rank must be >= top_n.")
    if start_date > as_of_date:
        raise ValueError("start_date cannot be later than as_of_date.")

    output_path = Path(excel_output)
    sheets = {} if reset_workbook else _load_workbook_sheets(output_path)
    metadata = _read_metadata(sheets)
    _check_settings(metadata, start_date, top_n, exit_rank, initial_capital)
    last_rebalance = metadata.get("Last_Rebalance_Date")
    cash_balance = float(metadata.get("Cash_Balance", initial_capital))
    positions = _position_map(sheets.get("Current_Portfolio"))
    summary_df = sheets.get("Monthly_Summary", pd.DataFrame(columns=SUMMARY_COLUMNS))
    rank_history_df = sheets.get("Rank_History", pd.DataFrame(columns=RANK_COLUMNS))
    trade_ledger_df = sheets.get("Trade_Ledger", pd.DataFrame(columns=TRADE_COLUMNS))

    scheduled_dates = monthly_rebalance_dates(
        start_date, as_of_date, date.fromisoformat(last_rebalance) if last_rebalance else None
    )
    if not scheduled_dates:
        portfolio_df = sheets.get(
            "Current_Portfolio", pd.DataFrame(columns=PORTFOLIO_COLUMNS)
        )
        sheets["Monthly_Summary"] = _enrich_monthly_summary(
            summary_df, initial_capital
        )
        sheets["Buy_Sell_Ledger"] = _build_buy_sell_ledger(
            trade_ledger_df.to_dict("records"), portfolio_df
        )
        if sheets:
            _write_workbook(output_path, sheets)
        return {
            "workbook": str(output_path.resolve()),
            "processed_dates": [],
            "current_holdings": len(positions),
            "cash_balance": cash_balance,
            "portfolio_value": (
                float(summary_df.iloc[-1]["Closing_Portfolio_Value"])
                if not summary_df.empty
                else cash_balance + sum(
                    float(position["Quantity"]) * float(position["Last_Price"])
                    for position in positions.values()
                )
            ),
        }

    panel, index_panel = load_cache_symbols(db_path)
    price_lookup = {
        symbol: frame["Close"].dropna()
        for symbol, frame in panel.items()
        if "Close" in frame.columns and not frame.empty
    }
    benchmark = index_panel.get("IDX:NIFTY 500")
    if benchmark is None or benchmark.empty:
        raise ValueError("The cache has no IDX:NIFTY 500 benchmark history.")

    summary_rows = summary_df.to_dict("records")
    rank_rows = rank_history_df.to_dict("records")
    trade_rows = trade_ledger_df.to_dict("records")
    processed_dates = []
    previous_ranks = _previous_ranks(rank_history_df, last_rebalance)
    previous_close = (
        float(summary_rows[-1]["Closing_Portfolio_Value"])
        if summary_rows
        else initial_capital
    )

    for rebalance_date in scheduled_dates:
        available_benchmark = benchmark.loc[
            :pd.Timestamp(rebalance_date), "Close"
        ].dropna()
        if available_benchmark.empty:
            raise ValueError(
                f"No Nifty 500 close is available on or before {rebalance_date}."
            )
        execution_date = available_benchmark.index[-1].date()
        ranked = compute_ranked_stock_details(panel, index_panel, execution_date)
        if not ranked:
            raise ValueError(f"No stocks could be ranked for {execution_date}.")

        is_initial = last_rebalance is None
        positions, cash_balance, trades, statuses = _run_rebalance(
            ranked,
            positions,
            cash_balance,
            initial_capital,
            top_n,
            exit_rank,
            rebalance_date,
            execution_date,
            price_lookup,
            is_initial,
        )
        for row in _build_rank_rows(
            ranked, statuses, previous_ranks, rebalance_date, execution_date
        ):
            rank_rows.append(row)

        portfolio_df = _portfolio_frame(positions, price_lookup, execution_date)
        closing_value = float(portfolio_df["Market_Value"].sum()) + cash_balance
        realized_pnl = sum(float(trade["Realized_PnL"]) for trade in trades)
        monthly_pnl = closing_value - previous_close
        summary_rows.append({
            "Rebalance_Date": rebalance_date.isoformat(),
            "Execution_Date": execution_date.isoformat(),
            "Ranked_Stocks": len(ranked),
            "Holdings_Count": len(positions),
            "Buys": sum(trade["Action"] == "BUY" for trade in trades),
            "Sells": sum(trade["Action"] == "SELL" for trade in trades),
            "Realized_PnL": realized_pnl,
            "Opening_Portfolio_Value": previous_close,
            "Closing_Portfolio_Value": closing_value,
            "Monthly_Profit_Loss": monthly_pnl,
            "Monthly_Gain_%": (
                monthly_pnl / previous_close * 100 if previous_close else np.nan
            ),
            "Cumulative_Profit_Loss": closing_value - initial_capital,
            "Cumulative_Return_%": (closing_value / initial_capital - 1) * 100,
            "CAGR_%": _calculate_cagr(
                initial_capital,
                closing_value,
                str(summary_rows[0]["Execution_Date"])
                if summary_rows
                else execution_date.isoformat(),
                execution_date.isoformat(),
            ),
            "Cash_Balance": cash_balance,
        })
        trade_rows.extend(trades)
        processed_dates.append(rebalance_date.isoformat())
        previous_close = closing_value
        last_rebalance = rebalance_date.isoformat()
        previous_ranks = {row["Symbol"]: int(row["Rank"]) for row in ranked}

    metadata.update({
        "Start_Date": start_date.isoformat(),
        "Top_N": str(top_n),
        "Exit_Rank": str(exit_rank),
        "Initial_Capital": f"{initial_capital:.2f}",
        "Last_Rebalance_Date": last_rebalance or "",
        "Cash_Balance": f"{cash_balance:.2f}",
        "Ranking_Method": "RRG composite from portfolio_strategy_backtest_rank.py",
        "Sizing_Rule": "Equal initial allocation; retained quantities unchanged; replacements funded by sale proceeds only",
        "Profit_Basis": "Gross realized P&L before fees, taxes, and slippage",
    })
    portfolio_df = _portfolio_frame(positions, price_lookup, as_of_date)
    summary_df = _enrich_monthly_summary(
        pd.DataFrame(summary_rows, columns=SUMMARY_COLUMNS), initial_capital
    )
    guide_df = pd.DataFrame([
        {"Topic": "Strategy", "Description": f"Top {top_n} RRG-ranked stocks; monthly review on the 25th."},
        {"Topic": "Exit rule", "Description": f"Sell owned stocks ranked below {exit_rank}; retain ranks 1-{exit_rank}."},
        {"Topic": "Replacement rule", "Description": f"Buy the highest-ranked unowned stocks within ranks 1-{exit_rank}, using sale proceeds only."},
        {"Topic": "Initial sizing", "Description": "Equal rupee budget per stock; whole shares only; residual cash is retained."},
        {"Topic": "Returns", "Description": "Trade ledger P&L is gross and excludes fees, taxes, and slippage."},
        {"Topic": "Monthly use", "Description": "Keep using this workbook as --excel-output to preserve the portfolio and ledger."},
        {"Topic": "Weekend/holiday dates", "Description": "Use the latest cached Nifty 500 close on or before the scheduled 25th."},
    ])
    updated_sheets = {
        "Guide": guide_df,
        "Metadata": _metadata_frame(metadata),
        "Current_Portfolio": portfolio_df,
        "Monthly_Summary": summary_df,
        "Rank_History": pd.DataFrame(rank_rows, columns=RANK_COLUMNS),
        "Trade_Ledger": pd.DataFrame(trade_rows, columns=TRADE_COLUMNS),
        "Buy_Sell_Ledger": _build_buy_sell_ledger(trade_rows, portfolio_df),
    }
    _write_workbook(output_path, updated_sheets)
    return {
        "workbook": str(output_path.resolve()),
        "processed_dates": processed_dates,
        "current_holdings": len(positions),
        "cash_balance": cash_balance,
        "portfolio_value": previous_close,
    }


def main() -> None:
    args = parse_args()
    result = run_tracker(
        db_path=args.db_path,
        start_date=args.start_date,
        as_of_date=args.as_of_date,
        top_n=args.top_n,
        exit_rank=args.exit_rank,
        initial_capital=args.initial_capital,
        excel_output=args.excel_output,
        reset_workbook=args.reset_workbook,
    )
    if result["processed_dates"]:
        print(f"Processed rebalance dates: {', '.join(result['processed_dates'])}")
    else:
        print("No new 25th-of-month rebalance date is due.")
    print(f"Current holdings: {result['current_holdings']}")
    print(f"Cash balance: INR {result.get('cash_balance', 0.0):,.2f}")
    print(f"Portfolio value: INR {result.get('portfolio_value', 0.0):,.2f}")
    print(f"Excel workbook: {result['workbook']}")


if __name__ == "__main__":
    main()
